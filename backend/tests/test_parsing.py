"""Smoke tests for document parsing."""
from __future__ import annotations

from io import BytesIO

from docx import Document as DocxDocument

from app.services.parsing import parse_bytes


def _make_docx_bytes() -> bytes:
    doc = DocxDocument()
    doc.add_paragraph("ТЕХНИЧЕСКОЕ ЗАДАНИЕ")
    doc.add_paragraph("на поставку очистителей воздуха")
    table = doc.add_table(rows=2, cols=3)
    table.rows[0].cells[0].text = "№"
    table.rows[0].cells[1].text = "Наименование"
    table.rows[0].cells[2].text = "Кол-во"
    table.rows[1].cells[0].text = "1"
    table.rows[1].cells[1].text = "Очиститель воздуха УФ"
    table.rows[1].cells[2].text = "40"
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_parse_docx_bytes_extracts_paragraphs_and_tables() -> None:
    data = _make_docx_bytes()
    text = parse_bytes("sample.docx", data)
    assert "ТЕХНИЧЕСКОЕ ЗАДАНИЕ" in text
    assert "очистителей воздуха" in text
    assert "[Таблица 1]" in text
    assert "Очиститель воздуха УФ" in text
    assert "40" in text


def test_parse_unknown_extension_falls_back_to_utf8() -> None:
    text = parse_bytes("plain.unknown", "ТЗ".encode("utf-8"))
    assert text == "ТЗ"
