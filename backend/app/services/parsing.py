"""Document parsing: DOCX / PDF / XLSX -> normalized markdown-like text."""
from __future__ import annotations

from io import BytesIO
from pathlib import Path

from docx import Document as DocxDocument
from openpyxl import load_workbook

from app.core.logging import logger


def parse_document(path: Path) -> str:
    """Detect file type and parse it into plain text with table structure preserved."""
    suffix = path.suffix.lower()
    if suffix == ".docx":
        return _parse_docx(path)
    if suffix == ".pdf":
        return _parse_pdf(path)
    if suffix in {".xlsx", ".xlsm"}:
        return _parse_xlsx(path)
    if suffix == ".txt":
        return path.read_text(encoding="utf-8", errors="replace")
    raise ValueError(f"Unsupported document type: {suffix}")


def _parse_docx(path: Path) -> str:
    doc = DocxDocument(str(path))
    return _render_docx(doc)


def _parse_pdf(path: Path) -> str:
    try:
        import pdfplumber
    except ImportError as exc:  # pragma: no cover - dependency declared
        raise RuntimeError("pdfplumber is required for PDF parsing") from exc

    blocks: list[str] = []
    with pdfplumber.open(str(path)) as pdf:
        for page_num, page in enumerate(pdf.pages, start=1):
            blocks.append(f"\n[Страница {page_num}]")
            text = page.extract_text() or ""
            if text.strip():
                blocks.append(text.strip())
            for ti, table in enumerate(page.extract_tables() or [], start=1):
                blocks.append(f"\n[Таблица {page_num}.{ti}]")
                for row in table:
                    cells = [(c or "").strip() for c in row]
                    if any(cells):
                        blocks.append(" || ".join(cells))
    return "\n".join(blocks)


def _parse_xlsx(path: Path) -> str:
    wb = load_workbook(str(path), data_only=True, read_only=True)
    blocks: list[str] = []
    for sheet in wb.worksheets:
        blocks.append(f"\n[Лист: {sheet.title}]")
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c).strip() if c is not None else "" for c in row]
            if any(cells):
                blocks.append(" || ".join(cells))
    return "\n".join(blocks)


def parse_bytes(filename: str, data: bytes) -> str:
    """Parse a file from raw bytes; the suffix from filename selects the parser."""
    suffix = Path(filename).suffix.lower()
    if suffix == ".docx":
        return _parse_docx_bytes(data)
    if suffix == ".pdf":
        return _parse_pdf_bytes(data)
    if suffix in {".xlsx", ".xlsm"}:
        return _parse_xlsx_bytes(data)
    if suffix == ".txt":
        return data.decode("utf-8", errors="replace")
    logger.warning("Unknown extension {} - attempting utf-8 decode", suffix)
    return data.decode("utf-8", errors="replace")


def _parse_docx_bytes(data: bytes) -> str:
    doc = DocxDocument(BytesIO(data))
    return _render_docx(doc)


def _render_docx(doc) -> str:
    """Render a python-docx Document to our normalized text format.

    Top-level tables become ``[Таблица N]`` blocks. Nested tables
    (tables embedded inside another table's cell) are emitted as
    ``[Таблица N.M]`` blocks after the top-level ones, and the
    enclosing parent cell text is augmented with a
    ``[Вложенная таблица N.M]`` reference so the extractor can
    associate them.
    """
    blocks: list[str] = []
    nested_blocks: list[str] = []

    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            blocks.append(text)

    for ti, table in enumerate(doc.tables, start=1):
        _render_table(table, str(ti), blocks, nested_blocks)

    blocks.extend(nested_blocks)
    return "\n".join(blocks)


def _render_table(
    table,
    table_id: str,
    out_blocks: list[str],
    nested_blocks: list[str],
) -> None:
    """Render one table (and recursively any nested ones) into out_blocks."""
    out_blocks.append(f"\n[Таблица {table_id}]")
    nested_counter = 0
    for row in table.rows:
        cells_text: list[str] = []
        seen_in_row: set[int] = set()
        for cell in row.cells:
            cell_key = id(cell._tc)
            if cell_key in seen_in_row:
                # Merged cell: python-docx repeats the same underlying _tc
                # across row.cells entries; only render it once.
                continue
            seen_in_row.add(cell_key)
            paragraphs_text = " | ".join(
                p.text.strip() for p in cell.paragraphs if p.text.strip()
            )
            cell_text = paragraphs_text
            for nested in cell.tables:
                nested_counter += 1
                nested_id = f"{table_id}.{nested_counter}"
                _render_table(nested, nested_id, nested_blocks, nested_blocks)
                ref = f"[Вложенная таблица {nested_id}]"
                cell_text = (cell_text + " " + ref).strip() if cell_text else ref
            cells_text.append(cell_text)
        if any(cells_text):
            out_blocks.append(" || ".join(cells_text))


def _parse_pdf_bytes(data: bytes) -> str:
    import pdfplumber

    blocks: list[str] = []
    with pdfplumber.open(BytesIO(data)) as pdf:
        for page_num, page in enumerate(pdf.pages, start=1):
            blocks.append(f"\n[Страница {page_num}]")
            text = page.extract_text() or ""
            if text.strip():
                blocks.append(text.strip())
            for ti, table in enumerate(page.extract_tables() or [], start=1):
                blocks.append(f"\n[Таблица {page_num}.{ti}]")
                for row in table:
                    cells = [(c or "").strip() for c in row]
                    if any(cells):
                        blocks.append(" || ".join(cells))
    return "\n".join(blocks)


def _parse_xlsx_bytes(data: bytes) -> str:
    wb = load_workbook(BytesIO(data), data_only=True, read_only=True)
    blocks: list[str] = []
    for sheet in wb.worksheets:
        blocks.append(f"\n[Лист: {sheet.title}]")
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c).strip() if c is not None else "" for c in row]
            if any(cells):
                blocks.append(" || ".join(cells))
    return "\n".join(blocks)
