"""Excel export of items and their supplier offers."""
from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.models.document import Document, Item


def build_xlsx(document: Document, items: list[Item]) -> bytes:
    wb = Workbook()

    ws_items = wb.active
    if ws_items is None:
        ws_items = wb.create_sheet("Позиции")
    ws_items.title = "Позиции"
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="2563EB")
    header_align = Alignment(horizontal="left", vertical="center", wrap_text=True)

    headers = ["№", "Наименование", "Кол-во", "Ед.изм.", "ГОСТ", "ОКПД2", "Характеристики"]
    ws_items.append(headers)
    for col_index, _ in enumerate(headers, start=1):
        cell = ws_items.cell(row=1, column=col_index)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align

    for idx, item in enumerate(items, start=1):
        specs = ""
        if item.specifications:
            specs = "\n".join(f"{k}: {v}" for k, v in item.specifications.items())
        ws_items.append(
            [
                idx,
                item.name,
                item.quantity,
                item.unit,
                item.gost,
                item.okpd2,
                specs,
            ]
        )

    _autosize(ws_items, [4, 60, 8, 10, 18, 18, 60])

    ws_offers = wb.create_sheet("Предложения")
    offer_headers = [
        "№ позиции",
        "Позиция",
        "Поставщик (домен)",
        "Заголовок",
        "Ссылка",
        "Цена",
        "Телефон",
        "Email",
        "Фрагмент",
    ]
    ws_offers.append(offer_headers)
    for col_index, _ in enumerate(offer_headers, start=1):
        cell = ws_offers.cell(row=1, column=col_index)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align

    for idx, item in enumerate(items, start=1):
        for offer in item.offers:
            ws_offers.append(
                [
                    idx,
                    item.name,
                    offer.supplier_domain,
                    offer.title,
                    offer.url,
                    offer.price,
                    offer.contact_phone,
                    offer.contact_email,
                    offer.snippet,
                ]
            )

    _autosize(ws_offers, [10, 40, 28, 50, 60, 14, 18, 28, 50])

    ws_matches = wb.create_sheet("Сопоставления")
    match_headers = [
        "№ позиции",
        "Позиция",
        "Гипотеза (марка / модель)",
        "Описание гипотезы",
        "Найденный товар",
        "Сайт",
        "Цена",
        "Совпадение, %",
        "Вердикт",
        "Сопоставление характеристик",
        "Резюме",
        "Ссылка",
    ]
    ws_matches.append(match_headers)
    for col_index, _ in enumerate(match_headers, start=1):
        cell = ws_matches.cell(row=1, column=col_index)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align

    for idx, item in enumerate(items, start=1):
        for match in item.matches:
            specs_text = ""
            if match.specs_compared:
                lines = []
                for s in match.specs_compared:
                    status = s.get("status", "unknown")
                    marker = {"match": "+", "mismatch": "-", "unknown": "?"}.get(status, "?")
                    lines.append(
                        f"[{marker}] {s.get('name', '')}: ТЗ={s.get('required', '—')} | "
                        f"найдено={s.get('found', '—')}"
                    )
                specs_text = "\n".join(lines)
            brand_model = " ".join(
                filter(None, [match.hypothesis_brand, match.hypothesis_model])
            ) or "—"
            ws_matches.append(
                [
                    idx,
                    item.name,
                    brand_model,
                    match.hypothesis_description,
                    match.found_title,
                    match.found_domain,
                    match.found_price,
                    match.match_score,
                    match.verdict,
                    specs_text,
                    match.summary,
                    match.found_url,
                ]
            )

    _autosize(ws_matches, [8, 36, 28, 36, 40, 22, 12, 14, 12, 60, 50, 60])

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _autosize(ws, widths: list[int]) -> None:
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[1].height = 32
